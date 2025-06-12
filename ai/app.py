from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from typing import List, Dict, Optional
import json
import pdfplumber
import os
from PIL import Image
import io
import ezdxf
import docx
import openpyxl
import pytesseract

app = FastAPI()

# Load checklist
with open('checklist.json', 'r') as f:
    CHECKLIST = json.load(f)

def extract_text_pdf(file):
    try:
        with pdfplumber.open(file) as pdf:
            return "\n".join(page.extract_text() or '' for page in pdf.pages)
    except Exception:
        # Try OCR for scanned PDFs
        file.seek(0)
        from pdf2image import convert_from_bytes
        images = convert_from_bytes(file.read())
        text = ""
        for img in images:
            text += pytesseract.image_to_string(img)
        return text

def extract_text_image(file):
    image = Image.open(file)
    return pytesseract.image_to_string(image)

def extract_text_dxf(file):
    try:
        doc = ezdxf.read(file)
        text = []
        for entity in doc.modelspace():
            if entity.dxftype() == 'TEXT':
                text.append(entity.dxf.text)
        return "\n".join(text)
    except Exception as e:
        return ""

def extract_text_docx(file):
    doc = docx.Document(file)
    return "\n".join([p.text for p in doc.paragraphs])

def extract_text_xlsx(file):
    wb = openpyxl.load_workbook(file)
    text = []
    for sheet in wb:
        for row in sheet.iter_rows(values_only=True):
            text.extend([str(cell) for cell in row if cell is not None])
    return "\n".join(text)

def extract_text_by_ext(file: UploadFile):
    ext = os.path.splitext(file.filename)[1].lower()
    if ext == '.pdf':
        return extract_text_pdf(file.file)
    elif ext in ['.jpg', '.jpeg', '.png', '.tiff']:
        return extract_text_image(file.file)
    elif ext == '.dxf':
        return extract_text_dxf(file.file)
    elif ext == '.dwg':
        raise HTTPException(status_code=400, detail='DWG files are not directly supported. Please convert to PDF or DXF before uploading.')
    elif ext == '.docx':
        return extract_text_docx(file.file)
    elif ext == '.xlsx':
        return extract_text_xlsx(file.file)
    else:
        raise HTTPException(status_code=400, detail=f'Unsupported file type: {ext}')

def checklist_status(extracted_text: str):
    missing_items = []
    present_items = []
    for rule in CHECKLIST:
        if rule['required'] and rule['description'].split(':')[0] not in extracted_text:
            missing_items.append(rule['description'])
        else:
            present_items.append(rule['description'])
    return missing_items, present_items

@app.post('/checklist-validate')
def checklist_validate(
    permitId: str = Form(...),
    currentFile: UploadFile = File(...),
    previousFile: Optional[UploadFile] = File(None)
):
    # Extract text from current file
    current_text = extract_text_by_ext(currentFile)
    current_missing, current_present = checklist_status(current_text)

    # If previous file is provided, do redline comparison
    if previousFile:
        previous_text = extract_text_by_ext(previousFile)
        previous_missing, previous_present = checklist_status(previous_text)
        # What was fixed: previously missing, now present
        fixed_items = [item for item in previous_missing if item in current_present]
        # What remains missing
        still_missing = current_missing
        # What is new (present in current, not in previous at all)
        new_items = [item for item in current_present if item not in previous_present]
        redline_changes = {
            "fixedItems": fixed_items,
            "stillMissing": still_missing,
            "newItems": new_items
        }
    else:
        redline_changes = None

    checklist_valid = len(current_missing) == 0

    return {
        "permitId": permitId,
        "checklistValid": checklist_valid,
        "missingItems": current_missing,
        "redlineChanges": redline_changes
    } 